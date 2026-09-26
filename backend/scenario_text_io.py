"""剧本文件读取：txt/md/pdf/docx/doc 文本提取与清洗。

从 backend.scenario_importer 拆出。
"""
from __future__ import annotations

import asyncio
import io
import json
import math
import os
import re
import subprocess
from typing import Any

from openai import AsyncOpenAI

from backend.config import ensure_valid_api_key, settings
from backend.engine.game_systems import detect_game_system
from backend.engine.llm_utils import strip_refusal as _strip_refusal
from backend.engine.prompt_guard import extract_json_array, extract_json_object, sanitize_user_text
from backend.engine.rag_utils import cosine as dense_cosine, embed_text
from backend.logging_utils import get_logger



# ═══════════════════════════════════════════════════════════════
# 文档读取
# ═══════════════════════════════════════════════════════════════

SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx", ".doc"}



def extract_text(filename: str, data: bytes) -> str:
    """从上传文件内容中提取纯文本。"""
    ext = os.path.splitext(filename)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的剧本格式: {ext or '(无扩展名)'}，支持: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")

    if ext in (".txt", ".md", ".markdown"):
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    if ext == ".doc":
        return _extract_doc(data)
    raise ValueError("无法读取该文件")



def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfStreamError
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                continue
        text = "\n\n".join(pages)
        if not text.strip():
            raise ValueError("PDF 中没有可提取的文本（可能是扫描件，请改用文字版 PDF）")
        return text
    except (PdfStreamError, Exception) as e:
        raise ValueError(f"PDF 解析失败：文件可能损坏或不是有效 PDF（{e}）") from e



def _extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if not text.strip():
        raise ValueError("DOCX 中没有可提取的文本")
    return text



def _extract_doc(data: bytes) -> str:
    """老式 .doc 是 OLE 复合格式，跨平台提取较麻烦。

    优先调用系统 antiword；不可用时退回提取可打印 ASCII/UTF-8 片段。
    """
    try:
        proc = subprocess.run(
            ["antiword", "-"],
            input=data,
            capture_output=True,
            timeout=20,
        )
        if proc.returncode == 0:
            text = proc.stdout.decode("utf-8", errors="replace").strip()
            if text:
                return text
    except Exception:
        pass

    # 退回：尝试 UTF-8 解码，失败则只保留可打印字符
    try:
        text = data.decode("utf-8", errors="strict")
    except UnicodeDecodeError:
        text = "".join(chr(b) for b in data if b in (9, 10, 13) or 32 <= b < 127)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError("无法从 .doc 中提取文本，请转换为 .docx / .txt / .pdf 后重试")
    return text



# ═══════════════════════════════════════════════════════════════
# 文本切分
# ═══════════════════════════════════════════════════════════════

def normalize_text(text: str) -> str:
    """清理文档中的常见噪声。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
